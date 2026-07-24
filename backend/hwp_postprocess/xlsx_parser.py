"""XLSX 파서 — openpyxl → 시트별 Markdown 테이블"""
import os
import shutil
from pathlib import Path

from openpyxl import load_workbook

from hwp_postprocess.models import ParseResult, PageResult, ParseStatus


def parse(file_path: str) -> ParseResult:
    source = str(file_path)
    # [수정 2026-07-25] .xls(옛 OLE2/BIFF 바이너리 포맷)는 openpyxl이 애초에 읽을 수
    # 없는 포맷이라(에러: "openpyxl does not support the old .xls file format") 지금까지
    # .xls 업로드가 전부 실패하고 있었다 — HWP를 LibreOffice로 DOCX 변환 후 읽는 것과
    # 같은 방식으로, .xls만 LibreOffice로 .xlsx 변환을 먼저 거치고 openpyxl로 읽는다.
    xlsx_path = file_path
    tmp_dir: str | None = None
    try:
        if Path(file_path).suffix.lower() == ".xls":
            from hwp_postprocess._libreoffice import convert
            xlsx_path = convert(file_path, "xlsx")
            tmp_dir = os.path.dirname(xlsx_path)

        # data_only=True: 수식 대신 캐시된 값을 읽는다
        wb = load_workbook(xlsx_path, data_only=True)
        pages: list[PageResult] = []

        for idx, sheet_name in enumerate(wb.sheetnames, start=1):
            ws = wb[sheet_name]
            md = _sheet_to_markdown(ws, sheet_name)
            pages.append(PageResult(page_num=idx, markdown=md))

        if not pages:
            return ParseResult(source_file=source, status=ParseStatus.FAILED, error="시트 없음")

        return ParseResult(source_file=source, status=ParseStatus.OK, pages=pages)
    except Exception as e:
        return ParseResult(source_file=source, status=ParseStatus.FAILED, error=str(e))
    finally:
        if tmp_dir:
            shutil.rmtree(tmp_dir, ignore_errors=True)


def _sheet_to_markdown(ws, sheet_name: str) -> str:
    rows: list[list[str]] = []

    for row in ws.iter_rows():
        cells = [str(cell.value) if cell.value is not None else "" for cell in row]
        # 완전히 빈 행 제외
        if any(c.strip() for c in cells):
            rows.append([c.replace("|", "\\|").replace("\n", " ") for c in cells])

    if not rows:
        return f"## {sheet_name}\n\n(빈 시트)"

    max_cols = max(len(r) for r in rows)
    rows = [r + [""] * (max_cols - len(r)) for r in rows]

    lines = [f"## {sheet_name}", ""]
    lines.append("| " + " | ".join(rows[0]) + " |")
    lines.append("| " + " | ".join(["---"] * max_cols) + " |")
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")

    return "\n".join(lines)

